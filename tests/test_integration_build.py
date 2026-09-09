from pathlib import Path
import importlib
import shutil
from zipfile import ZipFile

import pytest
from docx import Document
from docx.shared import Pt

from asc.compiler import build
from asc.findings import Status
from asc.models import dump_profile, load_profile


pytestmark = pytest.mark.skipif(shutil.which("pandoc") is None, reason="Pandoc is required for integration builds")


def _docx_xml(path: Path) -> str:
    with ZipFile(path) as archive:
        return "\n".join(archive.read(name).decode("utf-8", errors="ignore") for name in archive.namelist() if name.endswith(".xml"))


def _copy_demo_project(tmp_path: Path, root: Path, journal_id: str) -> tuple[Path, Path]:
    project = tmp_path / "中文 测试 project"
    for name in ("filters", "bibliography", "examples"):
        shutil.copytree(root / name, project / name)
    journal = project / "journals" / journal_id
    shutil.copytree(root / "journals/example-humanities-journal", journal)
    return project, journal


def test_real_pandoc_formatter_inspector_pipeline(root: Path):
    result = build(root / "examples/demo-paper/paper.md", "example-humanities-journal", root, live_zotero=False)
    assert result.docx.exists() and result.report.exists()
    assert result.final_status != Status.FAIL
    xml = _docx_xml(result.docx)
    assert "@hayles1999" not in xml
    document = Document(result.docx)
    assert document.paragraphs[0].text == "可复现的学术投稿工作流"
    assert any(p.style.name == "Abstract" for p in document.paragraphs)
    assert any(p.style.name == "Keywords" for p in document.paragraphs)
    assert any(p.style.name == "Bibliography" for p in document.paragraphs)
    with ZipFile(result.docx) as archive:
        assert "word/footnotes.xml" in archive.namelist()
    assert abs(document.sections[0].left_margin.mm - 31.7) < 0.2
    report = result.report.read_text(encoding="utf-8")
    assert "## DOCX Formatting Compliance" in report
    assert "## Profile Coverage" in report


def test_build_fails_final_status_when_formatter_outputs_wrong_size(monkeypatch, root: Path):
    module = importlib.import_module("asc.compiler.build")
    real_formatter = module.format_docx

    def wrong_formatter(input_path, output_path, profile):
        real_formatter(input_path, output_path, profile)
        document = Document(output_path)
        document.styles["Title"].font.size = Pt(7)
        document.save(output_path)

    monkeypatch.setattr(module, "format_docx", wrong_formatter)
    result = module.build(root / "examples/demo-paper/paper.md", "example-humanities-journal", root, live_zotero=False)
    assert result.final_status == Status.FAIL
    assert any(item.status == Status.FAIL and item.rule_path == "style.Title.size" for item in result.docx_findings)
    assert "Overall: **FAIL**" in result.report.read_text(encoding="utf-8")


def test_anonymous_build_respects_individual_hide_flags(tmp_path: Path, root: Path):
    project, journal = _copy_demo_project(tmp_path, root, "anonymous-test")
    profile = load_profile(journal / "profile.yaml")
    profile.journal.id = "anonymous-test"
    profile.anonymous_review.required = True
    profile.anonymous_review.hide_authors = True
    profile.anonymous_review.hide_affiliations = True
    profile.anonymous_review.hide_funding = False
    profile.anonymous_review.hide_email = True
    profile.anonymous_review.hide_orcid = True
    dump_profile(profile, journal / "profile.yaml")
    paper = project / "examples/demo-paper/paper.md"
    text = paper.read_text(encoding="utf-8")
    text = text.replace("affiliation: 示例大学数字人文实验室", "affiliation: 示例大学数字人文实验室\n    email: author@example.org\n    orcid: 0000-0002-1825-0097")
    paper.write_text(text, encoding="utf-8")
    result = build(paper, "anonymous-test", project, live_zotero=False)
    assert result.final_status != Status.FAIL
    xml = _docx_xml(result.docx)
    assert "示例作者" not in xml
    assert "示例大学数字人文实验室" not in xml
    assert "author@example.org" not in xml
    assert "0000-0002-1825-0097" not in xml
    assert "虚构演示项目" in xml
    assert not [item for item in result.docx_findings if item.status == Status.FAIL]


def test_custom_keyword_separators_reach_docx(tmp_path: Path, root: Path):
    from asc.models import SourceKind

    project, journal = _copy_demo_project(tmp_path, root, "separator-test")
    profile = load_profile(journal / "profile.yaml")
    profile.journal.id = "separator-test"
    profile.keywords.zh.separator = "，"
    profile.keywords.zh.separator_source_kind = SourceKind.user_override
    profile.keywords.en.separator = " | "
    profile.keywords.en.separator_source_kind = SourceKind.user_override
    dump_profile(profile, journal / "profile.yaml")
    result = build(project / "examples/demo-paper/paper.md", "separator-test", project, live_zotero=False)
    assert result.final_status != Status.FAIL
    text = "\n".join(paragraph.text for paragraph in Document(result.docx).paragraphs)
    assert "学术写作，可复现工作流，文档编译" in text
    assert "academic writing | reproducible workflow | document compilation" in text
    separator_findings = [item for item in result.docx_findings if item.rule_path and item.rule_path.endswith(".separator")]
    assert separator_findings and all(item.status == Status.PASS for item in separator_findings)
